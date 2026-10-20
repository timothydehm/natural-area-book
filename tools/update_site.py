#!/usr/bin/env python3
"""Rebuild the site's content from the master spreadsheet.

  * Reads the spreadsheet: the Google Sheet named by the SHEET_ID environment variable
    (shared "Anyone with the link can view"), or else the .xlsx in content/.
  * Writes data.js, which the site loads.
  * Makes web-sized copies of everything in photos/ into photos/web/.

Run by the "Update site" GitHub Action, or locally:  python tools/update_site.py
Needs: pip install openpyxl pillow
"""
import io, json, os, re, sys, urllib.request
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
from PIL import Image, ImageOps

ROOT = Path(__file__).resolve().parent.parent
PHOTOS, WEB = ROOT / "photos", ROOT / "photos" / "web"
WEB_LONG_SIDE = 1800
IMAGE_TYPES = {".jpg", ".jpeg", ".png", ".webp", ".heic", ".tif", ".tiff"}


def nums(s):
    return [int(n) for n in re.findall(r"\((\d+)\)", str(s or ""))]


def txt(v):
    return str(v).strip() if v not in (None, "") else None


def open_workbook():
    sheet_id = os.environ.get("SHEET_ID", "").strip()
    if sheet_id:
        url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx"
        print(f"Reading Google Sheet {sheet_id}")
        with urllib.request.urlopen(url, timeout=60) as r:
            body = r.read()
        if not body.startswith(b"PK"):
            sys.exit("Google didn't return a spreadsheet. Is the sheet shared as 'Anyone with the link can view'?")
        return openpyxl.load_workbook(io.BytesIO(body), data_only=True), "Google Sheet"
    files = sorted((ROOT / "content").glob("*.xlsx"))
    if not files:
        sys.exit("No SHEET_ID set and no .xlsx in content/.")
    print(f"Reading {files[0].name}")
    return openpyxl.load_workbook(files[0], data_only=True), files[0].name


def read_data(wb):
    rows = list(wb["Patterns"].iter_rows(values_only=True))
    head = [str(h or "").strip().lower() for h in rows[0]]
    col = lambda *names: next((head.index(n) for n in names if n in head), None)
    ix = {"num": col("#"), "part": col("part"), "name": col("pattern"), "see": col("what you see"),
          "why": col("why it matters"), "do": col("what to do"), "rc": col("related conditions"),
          "rp": col("related design patterns", "related patterns"), "photo": col("photo"),
          "crop": col("photo crop", "crop")}
    pats = []
    for r in rows[1:]:
        g = lambda k: (r[ix[k]] if ix[k] is not None and ix[k] < len(r) else None)
        if g("num") in (None, ""):
            continue
        pats.append({"num": int(g("num")), "part": 2 if "2" in str(g("part")) else 1, "name": txt(g("name")),
                     "see": txt(g("see")), "why": txt(g("why")), "do": txt(g("do")),
                     "rc": nums(g("rc")), "rp": nums(g("rp")), "photo": txt(g("photo")), "crop": txt(g("crop"))})
        # (the Notes column is deliberately left out: review notes stay private)
    s = {}
    if "Book" in wb.sheetnames:
        for r in list(wb["Book"].iter_rows(values_only=True))[1:]:
            if r and r[0]:
                s[str(r[0]).strip().lower()] = "" if r[1] is None else str(r[1]).strip()
    settings = {
        "title": s.get("title") or "Natural Areas Pattern Book", "kicker": s.get("kicker", "A field guide"),
        "edition": s.get("edition", ""), "intro": s.get("intro", ""), "howTo": s.get("how to use", ""),
        "coverPhoto": s.get("cover photo") or None, "coverCrop": s.get("cover photo crop") or None,
        "parts": {"1": {"name": s.get("part 1 name") or "Conditions", "intro": s.get("part 1 intro", "")},
                  "2": {"name": s.get("part 2 name") or "Design Patterns", "intro": s.get("part 2 intro", "")}},
    }
    return settings, sorted(pats, key=lambda p: p["num"])


def make_web_photos():
    WEB.mkdir(parents=True, exist_ok=True)
    out = {}
    for src in sorted(PHOTOS.iterdir()):
        if not src.is_file() or src.suffix.lower() not in IMAGE_TYPES:
            continue
        dst = WEB / (src.stem + ".jpg")
        out[src.name] = f"photos/web/{dst.name}"
        # locally, skip photos that are already converted; on GitHub (CI) always rebuild
        if not os.environ.get("CI") and dst.exists() and dst.stat().st_mtime >= src.stat().st_mtime:
            continue
        try:
            im = ImageOps.exif_transpose(Image.open(src)).convert("RGB")
            im.thumbnail((WEB_LONG_SIDE, WEB_LONG_SIDE), Image.LANCZOS)
            im.save(dst, "JPEG", quality=78, optimize=True, progressive=True)
            print(f"  photo: {src.name} -> {dst.relative_to(ROOT)}")
        except Exception as e:
            print(f"  WARNING: couldn't convert {src.name}: {e}")
            out.pop(src.name, None)
    return out


def main():
    wb, source = open_workbook()
    settings, pats = read_data(wb)
    photos = make_web_photos()
    known = {p["num"] for p in pats}
    for p in pats:
        for n in p["rc"] + p["rp"]:
            if n not in known:
                print(f"  WARNING: pattern {p['num']} refers to ({n}), which doesn't exist")
        if p["photo"] and p["photo"] not in photos:
            print(f"  WARNING: pattern {p['num']}: photo '{p['photo']}' is not in photos/")
    data = {"settings": settings, "patterns": pats, "photos": photos, "source": source,
            "updated": datetime.now(timezone.utc).strftime("%Y-%m-%d")}
    (ROOT / "data.js").write_text("window.BOOK_DATA = " + json.dumps(data, ensure_ascii=False, indent=1) + ";\n",
                                  encoding="utf-8")
    print(f"Wrote data.js: {len(pats)} patterns, {len(photos)} photos")


if __name__ == "__main__":
    main()
